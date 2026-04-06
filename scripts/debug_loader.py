import openpyxl

wb = openpyxl.load_workbook('data/sections.xlsx', data_only=True)
ws = wb.active
headers = [cell.value for cell in ws[1]]
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[0]: continue
    row_data = dict(zip(headers, row))
    print(row_data)
