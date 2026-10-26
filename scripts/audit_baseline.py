"""HFS-000: capture observed behavior, NOT an independent engineering oracle.

Run from the repository with a fresh audit environment. Output/work directories
must not exist: frozen evidence is never overwritten. No production files change.
"""
import argparse
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET


ROOT = Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def execute(command, cwd, env=None):
    completed = subprocess.run(command, cwd=cwd, env=env, capture_output=True,
                               text=True, encoding="utf-8", errors="replace", timeout=240)
    return {"command": list(map(str, command)), "cwd": str(cwd),
            "returncode": completed.returncode, "stdout": completed.stdout,
            "stderr": completed.stderr}


def normalize(value):
    """Remove only volatile timestamps/paths; keep all engineering numbers."""
    volatile = {"generated_at_utc", "input_file", "sections_file", "results_run_dir",
                "detailed_results_json", "interaction_diagram_i5_method2_image"}
    if isinstance(value, dict):
        return {key: (None if val is None else "<volatile>") if key in volatile
                else normalize(val) for key, val in value.items()}
    if isinstance(value, list):
        return [normalize(item) for item in value]
    return value


def canonical_hash(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"),
                                    ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def engine_snapshot(work):
    # Import exactly the byte-identical copy under test, not the working checkout.
    sys.path.insert(0, str(work))
    from src.steel_design.engine import SteelDesignEngine
    from src.steel_design.models import SteelCase
    from src.steel_design.run import _extract_output_options, _merge_case_payload
    from src.steel_design.sections.loader import load_sections_from_excel
    from src.steel_design.sections.repository import SectionRepository
    import openpyxl

    sections = load_sections_from_excel(str(work / "data/sections.xlsx"))
    repo = SectionRepository()
    repo.load(sections)
    names = Counter(s.name.upper() for s in sections)
    workbook = openpyxl.load_workbook(work / "data/sections.xlsx", read_only=True, data_only=True)
    catalog = {"loaded": len(sections), "repository_count": len(repo.get_all()),
               "families": dict(sorted(Counter(s.type for s in sections).items())),
               "duplicates": {name: n for name, n in names.items() if n > 1},
               "sheets": {s.title: {"rows": s.max_row, "columns": s.max_column} for s in workbook}}
    workbook.close()
    engine = SteelDesignEngine(repo)
    outputs = {}
    for path in sorted((work / "examples").glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        base = {k: v for k, v in payload.items() if k != "cases"}
        cases = payload.get("cases", [{}])
        results = []
        for i, override in enumerate(cases):
            merged = _merge_case_payload(base, {k: v for k, v in override.items() if k != "name"})
            cleaned, _ = _extract_output_options(merged)
            try:
                result = engine.run(SteelCase(**cleaned)).model_dump(mode="json")
            except Exception as exc:
                result = {"error": {"type": type(exc).__name__, "message": str(exc),
                                    "code": getattr(exc, "code", None), "details": getattr(exc, "details", None)}}
            results.append({"case": override.get("name", f"case_{i+1}"), "result": result})
        outputs[path.name] = results
    return catalog, outputs


def capture(args):
    out, work = Path(args.out).resolve(), Path(args.work).resolve()
    if out.exists() or work.exists():
        raise SystemExit("Use new output and work directories; existing evidence is immutable.")
    out.mkdir(parents=True)
    work.mkdir(parents=True)
    git = args.git
    tracked_run = execute([git, "ls-files", "-z"], ROOT)
    if tracked_run["returncode"]:
        raise RuntimeError(tracked_run)
    tracked = [p for p in tracked_run["stdout"].split("\x00") if p]
    inventory = {}
    for name in tracked:
        path = ROOT / name
        try:
            inventory[name] = {"bytes": path.stat().st_size, "sha256": sha(path)}
        except OSError as exc:
            inventory[name] = {"read_error": str(exc)}
    write_json(out / "tracked_inventory.json", inventory)
    write_json(out / "git_before.json", {
        "head": execute([git, "rev-parse", "HEAD"], ROOT),
        "branch": execute([git, "branch", "--show-current"], ROOT),
        "status": execute([git, "status", "--short", "--untracked-files=all"], ROOT),
        "diff": execute([git, "diff", "--stat", "HEAD"], ROOT),
        "version": execute([git, "--version"], ROOT)})
    copied = {}
    paths = sorted(set([p for p in (ROOT / "src").rglob("*.py") if "__pycache__" not in p.parts]
                       + list((ROOT / "tests").glob("test_*.py"))
                       + list((ROOT / "examples").glob("*.json"))
                       + [ROOT / "data/sections.xlsx", ROOT / "requirements.txt"]))
    for source in paths:
        relative = source.relative_to(ROOT)
        target = work / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        assert sha(source) == sha(target)
        copied[relative.as_posix()] = sha(target)
    write_json(out / "copied_inputs.json", copied)
    freeze = execute([sys.executable, "-m", "pip", "freeze", "--all"], ROOT)
    if freeze["returncode"]:
        raise RuntimeError(freeze)
    (out / "environment-freeze.txt").write_text(freeze["stdout"], encoding="utf-8")
    env = os.environ.copy()
    env.pop("PYTHONPATH", None)
    env["PYTHONNOUSERSITE"] = "1"
    env["PYTEST_DISABLE_PLUGIN_AUTOLOAD"] = "1"
    env["PYTHONUTF8"] = "1"
    write_json(out / "environment.json", {"python": sys.version, "executable": sys.executable,
        "platform": platform.platform(), "system_site_packages": False,
        "test_environment_overrides": {key: env[key] for key in
            ("PYTHONNOUSERSITE", "PYTEST_DISABLE_PLUGIN_AUTOLOAD", "PYTHONUTF8")},
        "pyvenv_cfg": (Path(sys.prefix) / "pyvenv.cfg").read_text(encoding="utf-8")})
    suite_codes = {}
    for label, extra in [("existing", ["--ignore=tests/test_hfs000_characterization.py"]),
                         ("characterization", ["tests/test_hfs000_characterization.py"])]:
        xml_path = out / f"pytest-{label}.xml"
        command = [sys.executable, "-m", "pytest", "-q", "-rA", "--tb=short", *extra,
                   f"--junitxml={xml_path}", "-o", "junit_logging=all"]
        run = execute(command, work, env)
        suite_codes[label] = run["returncode"]
        write_json(out / f"pytest-{label}.json", run)
        if xml_path.exists():
            print(label, [s.attrib for s in ET.parse(xml_path).getroot()], flush=True)
        else:
            print(label, run, flush=True)
    runs = []
    for repeat in (1, 2):
        records = {}
        for case in sorted((work / "examples").glob("*.json")):
            command = [sys.executable, "-m", "src.steel_design.run", f"examples/{case.name}",
                       "--results-dir", f"audit-results/repeat-{repeat}"]
            run = execute(command, work, env)
            stream = run["stderr"] if run["returncode"] else run["stdout"]
            try:
                payload = normalize(json.loads(stream))
            except json.JSONDecodeError:
                payload = {"unparsed_stdout": run["stdout"], "unparsed_stderr": run["stderr"]}
            records[case.name] = {"exit_code": run["returncode"], "payload": payload,
                                  "sha256_canonical": canonical_hash(payload)}
            write_json(out / "cli" / f"{case.stem}.repeat-{repeat}.json", run)
        runs.append(records)
    catalog, outputs = engine_snapshot(work)
    _, outputs_repeat = engine_snapshot(work)
    write_json(out / "catalog.json", catalog)
    write_json(out / "engine-results.json", outputs)
    write_json(out / "examples.json", runs[0])
    artifacts = {}
    for path in sorted((work / "audit-results").rglob("*")):
        if path.is_file():
            artifacts[path.relative_to(work).as_posix()] = {"bytes": path.stat().st_size, "sha256": sha(path)}
            target = out / "artifacts" / path.relative_to(work / "audit-results")
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, target)
    write_json(out / "artifacts.json", artifacts)
    changed = [name for name, record in inventory.items()
               if "sha256" in record and sha(ROOT / name) != record["sha256"]]
    verification = {"repeat_cli_identical_after_normalization": runs[0] == runs[1],
                    "repeat_engine_identical": outputs == outputs_repeat,
                    "pytest_exit_codes": suite_codes,
                    "original_tracked_files_changed": changed,
                    "engine_results_sha256": canonical_hash(outputs)}
    write_json(out / "verification.json", verification)
    write_json(out / "git_after.json", execute([git, "status", "--short", "--untracked-files=all"], ROOT))
    write_json(out / "SHA256SUMS.json", {p.relative_to(out).as_posix(): sha(p)
                                        for p in sorted(out.rglob("*")) if p.is_file()})
    print(json.dumps(verification, indent=2), flush=True)
    return 0 if (not changed and runs[0] == runs[1] and outputs == outputs_repeat
                 and all(code == 0 for code in suite_codes.values())) else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--git", default="git")
    parser.add_argument("--out", required=True)
    parser.add_argument("--work", required=True)
    raise SystemExit(capture(parser.parse_args()))
