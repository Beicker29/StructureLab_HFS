import shutil
import time
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path

import openpyxl

from src.steel_design.domain.errors import CatalogSchemaError
from src.steel_design.sections.loader import load_sections_from_excel


@contextmanager
def _temporary_test_dir():
    root = Path(__file__).resolve().parent / ".tmp_tests"
    root.mkdir(exist_ok=True)
    tmp_path = root / f"loader_{uuid.uuid4().hex}"
    tmp_path.mkdir(parents=True, exist_ok=False)
    try:
        yield tmp_path
    finally:
        for _ in range(5):
            try:
                shutil.rmtree(tmp_path)
                break
            except FileNotFoundError:
                break
            except PermissionError:
                time.sleep(0.2)
        else:
            shutil.rmtree(tmp_path, ignore_errors=True)


def _write_excel(path: Path, headers, rows):
    workbook = openpyxl.Workbook()
    try:
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
    finally:
        workbook.close()


class CatalogLoaderTests(unittest.TestCase):
    def test_catalog_loader_supports_aliases_and_profile_classification(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "area [mm2]", "Zx [mm3]", "rx [mm]", "ry [mm]"],
                rows=[
                    ["PTE 150X300X8.0mm", 6332.0, 120000.0, 91.0, 51.0],
                    ["PTE 200X8.0mm", 4800.0, 85000.0, 75.0, 75.0],
                    ["L6X6X1_2", 2500.0, 45000.0, 22.0, 22.0],
                ],
            )

            sections = load_sections_from_excel(str(excel_path))

            self.assertEqual(3, len(sections))
            self.assertEqual("RECTANGULAR_HSS", sections[0].type)
            self.assertEqual("CIRCULAR_HSS", sections[1].type)
            self.assertEqual("SINGLE_ANGLE", sections[2].type)

    def test_catalog_loader_fails_when_required_headers_are_missing(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "Zx [mm3]"],
                rows=[["W14X90", 123000.0]],
            )

            with self.assertRaises(CatalogSchemaError) as raised:
                load_sections_from_excel(str(excel_path))

            self.assertEqual("CATALOG_SCHEMA_ERROR", raised.exception.code)

    def test_catalog_loader_fails_for_unknown_profile_name(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "area [mm2]", "Zx [mm3]"],
                rows=[["UNKNOWN_PROFILE", 1000.0, 100.0]],
            )

            with self.assertRaises(CatalogSchemaError) as raised:
                load_sections_from_excel(str(excel_path))

            self.assertIn("shape", raised.exception.details)

    def test_catalog_loader_scales_units_from_aisc_readme_conventions(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "A", "Ix", "Zx", "J", "Cw", "rx", "ry"],
                rows=[["W14X90", 17000.0, 100.0, 200.0, 3.0, 4.0, 90.0, 50.0]],
            )

            sections = load_sections_from_excel(str(excel_path))
            section = sections[0]

            self.assertAlmostEqual(100.0e6, section.ix)
            self.assertAlmostEqual(200.0e3, section.zx)
            self.assertAlmostEqual(3.0e3, section.j)
            self.assertAlmostEqual(4.0e9, section.cw)

    def test_catalog_loader_classifies_channel_profile(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "A", "Ix", "Iy", "rx", "ry"],
                rows=[["C15X50", 9480.0, 168.0, 4.58, 133.0, 22.0]],
            )

            sections = load_sections_from_excel(str(excel_path))
            self.assertEqual("CHANNEL", sections[0].type)

    def test_catalog_loader_classifies_ipe_profile_as_w_family(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            workbook = openpyxl.Workbook()
            try:
                default_sheet = workbook.active
                workbook.remove(default_sheet)
                sheet = workbook.create_sheet("W_IPE_HEA_HEB")
                sheet.append(["shape", "A", "rx", "ry"])
                sheet.append(["IPE 80", 764.0, 31.0, 11.0])
                workbook.save(excel_path)
            finally:
                workbook.close()

            sections = load_sections_from_excel(str(excel_path))
            self.assertEqual(1, len(sections))
            self.assertEqual("W", sections[0].type)

    def test_catalog_loader_derives_slenderness_ratios_when_missing(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "A", "b", "h", "tdes", "ID", "rx", "ry", "J"],
                rows=[["PTE 100X80X4.0", 2000.0, 90.0, 70.0, 3.6, 90.0, 20.0, 15.0, 30.0]],
            )

            section = load_sections_from_excel(str(excel_path))[0]
            self.assertAlmostEqual(25.0, section.b_tdes_ratio, places=3)
            self.assertAlmostEqual(19.444444, section.h_tdes_ratio, places=3)
            self.assertAlmostEqual(25.0, section.d_t_ratio, places=3)

    def test_catalog_loader_treats_excel_error_tokens_as_missing_optional_values(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "A", "Sx", "rx", "ry"],
                rows=[["W14X90", 17000.0, "#DIV/0!", 90.0, 50.0]],
            )

            sections = load_sections_from_excel(str(excel_path))
            self.assertEqual(1, len(sections))
            self.assertIsNone(sections[0].sx)

    def test_catalog_loader_ignores_wt_profiles(self):
        with _temporary_test_dir() as tmp:
            excel_path = tmp / "sections.xlsx"
            _write_excel(
                excel_path,
                headers=["shape", "A", "rx", "ry"],
                rows=[
                    ["WT22X204", 15000.0, 90.0, 55.0],
                    ["W14X90", 17000.0, 90.0, 55.0],
                ],
            )

            sections = load_sections_from_excel(str(excel_path))
            self.assertEqual(1, len(sections))
            self.assertEqual("W14X90", sections[0].name)
            self.assertEqual("W", sections[0].type)


if __name__ == "__main__":
    unittest.main()
