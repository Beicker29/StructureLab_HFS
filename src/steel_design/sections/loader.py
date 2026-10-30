"""Read data/sections.xlsx (AISC v16 metric database layout plus PTE sheets) into Section records.

Strict: an unknown sheet, a missing required column or a non-numeric value raises
CatalogSchemaError. Placeholders ("–", "#DIV/0!", ...) mean "no data" and stay None.
"""
import hashlib
from io import BytesIO
from pathlib import Path

import openpyxl
from pydantic import TypeAdapter, ValidationError as PydanticValidationError

from ..errors import CatalogSchemaError
from ..models import FiniteNumber
from .identity import section_identity
from .models import Section, SectionSource

_NUMBER = TypeAdapter(FiniteNumber)

SHEET_FAMILY = {"W_IPE_HEA_HEB": "W", "C": "CHANNEL", "L": "SINGLE_ANGLE", "2L": "DOUBLE_ANGLE",
                "PTE_RECT": "RECTANGULAR_HSS", "PTE_RED": "CIRCULAR_HSS"}

# Section field -> catalog header (only properties used by Chapters D, E and F).
COLUMNS = {
    "area": "A", "d": "d", "h": "h", "bf": "bf", "b": "b", "od": "OD", "tf": "tf", "t": "t", "tdes": "tdes",
    "tw": "tw", "ho": "ho", "h_out": "Ht", "b_out": "B", "ix": "Ix", "iy": "Iy", "iw": "Iw", "iz": "Iz", "x": "x", "y": "y",
    "tan_alpha": "tan(α)", "sx": "Sx", "sy": "Sy", "zx": "Zx", "zy": "Zy",
    "sw_a": "SwA", "sw_b": "SwB", "sw_c": "SwC", "sz_a": "SzA", "sz_b": "SzB", "sz_c": "SzC", "rx": "rx", "ry": "ry", "rz": "rz", "j": "J", "cw": "Cw", "ro": "ro", "H": "H",
    "bf_2tf_ratio": "bf/2tf", "b_t_ratio": "b/t", "b_tdes_ratio": "b/tdes", "h_tw_ratio": "h/tw",
    "h_tdes_ratio": "h/tdes", "d_t_ratio": "D/t",
}

# AISC v16 metric database units: 10^6 mm4 (Ix, Iy, Iw, Iz), 10^3 mm4 (J), 10^3 mm3 (S, Z), 10^9 mm6 (Cw).
SCALE = {"ix": 1e6, "iy": 1e6, "iw": 1e6, "iz": 1e6, "j": 1e3, "cw": 1e9, "sx": 1e3, "sy": 1e3, "zx": 1e3, "zy": 1e3,
         "sw_a": 1e3, "sw_b": 1e3, "sw_c": 1e3, "sz_a": 1e3, "sz_b": 1e3, "sz_c": 1e3}
# Sheet-specific exceptions, verified with S = I/c (tests/test_catalog.py): PTE_RED stores Sx, Sy in mm3.
SHEET_SCALE = {"PTE_RED": {"sx": 1.0, "sy": 1.0}}

PLACEHOLDERS = {"", "-", "–", "—", "N/A", "NA", "NONE", "NULL",
                "#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NUM!", "#NAME?", "#NULL!"}


def _number(value, *, section, field, row, scale):
    if value is None or (isinstance(value, str) and value.strip().upper() in PLACEHOLDERS):
        return None
    try:
        if isinstance(value, bool):
            raise ValueError
        return _NUMBER.validate_python(float(value) * scale)
    except (TypeError, ValueError, PydanticValidationError):
        raise CatalogSchemaError("El valor del catálogo debe ser un número finito.",
                                 details={"section": section, "field": field, "value": repr(value), "row": row})


def _derive_ratios(values: dict, family: str) -> None:
    """B4.1 ratios the PTE sheets do not tabulate: flat width / tdes, OD / tdes (B4.1b(d), (g))."""
    tdes = values["tdes"]
    if not tdes:
        return
    if family == "RECTANGULAR_HSS":
        for ratio, width in (("b_tdes_ratio", "b"), ("h_tdes_ratio", "h")):
            if values[ratio] is None and values[width] is not None:
                values[ratio] = values[width] / tdes
    if family == "CIRCULAR_HSS" and values["d_t_ratio"] is None and values["od"] is not None:
        values["d_t_ratio"] = values["od"] / tdes


def load_sections_from_excel(filepath: str) -> list[Section]:
    # Hash exactly the bytes parsed, not a second read of a changing file.
    source_bytes = Path(filepath).read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    workbook = openpyxl.load_workbook(BytesIO(source_bytes), data_only=True)
    sections = []
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().upper() == "README":
            continue
        family = SHEET_FAMILY.get(sheet_name.strip())
        if family is None:
            raise CatalogSchemaError("Hoja del catálogo desconocida.", details={"sheet": sheet_name,
                                                                      "expected": sorted(SHEET_FAMILY)})
        sheet = workbook[sheet_name]
        headers = [str(h).strip() if h is not None else None for h in next(sheet.iter_rows(max_row=1, values_only=True))]
        missing = [h for h in ("shape", "A") if h not in headers]
        if missing:
            raise CatalogSchemaError("Faltan columnas obligatorias en la hoja del catálogo.",
                                     details={"sheet": sheet_name, "missing": missing})
        index = {field: headers.index(header) for field, header in COLUMNS.items() if header in headers}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            original = row[headers.index("shape")]
            if original is None or not str(original).strip():
                continue
            name = section_identity(str(original))
            scales = {**SCALE, **SHEET_SCALE.get(sheet_name.strip(), {})}
            values = {field: _number(row[i], section=name, field=field, row=row_number,
                                     scale=scales.get(field, 1.0)) if i is not None else None
                      for field, i in ((f, index.get(f)) for f in COLUMNS)}
            if values["area"] is None:
                raise CatalogSchemaError("Falta el área A, obligatoria.", details={"section": name, "row": row_number})
            _derive_ratios(values, family)
            try:
                sections.append(Section(
                    name=name, type=family,
                    source=SectionSource(file=str(Path(filepath).resolve()), sha256=source_hash,
                                         sheet=sheet_name, row=row_number, original_name=str(original)),
                    **values))
            except PydanticValidationError as exc:
                raise CatalogSchemaError("Propiedades del catálogo no válidas.", details={
                    "section": name, "sheet": sheet_name, "row": row_number,
                    "fields": exc.errors(include_url=False, include_context=False, include_input=False)}) from exc
    return sections
